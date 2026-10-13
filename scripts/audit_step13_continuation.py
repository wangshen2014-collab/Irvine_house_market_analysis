#!/usr/bin/env python3
"""Targeted Step 13 checks; never reads A/C or imports production modules.

Python 3 + openpyxl (read-only workbook extraction). Raw evidence stays private.
This is evidence/definition verification, not a market-return or timing test.
"""
import argparse
import csv
import hashlib
import json
import math
import sqlite3
import subprocess
import warnings
from collections import Counter
from decimal import Decimal, getcontext
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
FREEZE = '697f152000da23ad842be849d91a82ea484baa1d'
PRIOR = 'c2ad6899d8a3aa5a41fd36afa54e3b9a7f76e716'
getcontext().prec = 40


def rows(path):
    with path.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def num(v):
    try:
        d = float(v)
        return d if math.isfinite(d) else None
    except (ValueError, TypeError):
        return None


def pct(a, b):
    return float((Decimal(str(a)) / Decimal(str(b)) - 1) * 100)


def git_bytes(ref, path):
    return subprocess.check_output(['git', 'show', ref + ':' + path], cwd=ROOT)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--evidence-root', type=Path, required=True)
    ap.add_argument('--out', type=Path, required=True)
    args = ap.parse_args()
    out = {'reviewed_commit': FREEZE, 'household_A_C_read_or_run': False,
           'production_modules_imported': False, 'scope': 'targeted_new_step12_evidence'}
    index = json.loads(git_bytes(FREEZE, 'data/step12_evidence_index.json'))
    fingerprints = []
    for e in index['entries']:
        for record in e.get('originals', []):
            matches = list(args.evidence_root.rglob(Path(record['file']).name))
            assert len(matches) == 1, record['file']
            h = hashlib.sha256(matches[0].read_bytes()).hexdigest()
            assert h == record['sha256'], record['file']
            fingerprints.append({'evidence': e['id'], 'file': matches[0].name,
                                 'sha256': h, 'matches_frozen_index': True})
    out['raw_fingerprints'] = fingerprints
    acs = []
    for row in rows(ROOT/'data/clean/acs_household_window_changes.csv'):
        if row['geography'] != 'Irvine' or row['metric'] not in {'households_total','owner_households','renter_households'}:
            continue
        change = Decimal(row['new_estimate']) - Decimal(row['old_estimate'])
        moe = (Decimal(row['new_moe90'])**2 + Decimal(row['old_moe90'])**2).sqrt()
        assert change == Decimal(row['change'])
        assert abs(moe - Decimal(row['change_moe90'])) < Decimal('0.000001')
        acs.append({'metric':row['metric'],'change':float(change),'derived_moe90':float(moe),
                    'method':row['moe_method'],'raw_original_reread':False})
    out['acs_standard_table_derivation'] = acs
    inherited = ['scripts/audit_step13_financing.py', 'data/scenario_analysis_design.json',
                 'data/clean/scenario_equal_payment_boundaries.csv', 'data/clean/scenario_points.csv',
                 'data/clean/scenario_arm_cap_stress.csv', 'data/clean/scenario_refinance.csv',
                 'qa/step13_financing_independent.json']
    out['inherited_financing'] = []
    for p in inherited:
        old, new = git_bytes(PRIOR, p), git_bytes(FREEZE, p)
        assert old == new
        out['inherited_financing'].append({'path': p, 'sha256': hashlib.sha256(new).hexdigest(),
                                          'unchanged': True, 'rerun': False})
    # An independent SQL filter, numeric conversion and order-statistic median.
    db = sqlite3.connect(':memory:')
    raw = rows(args.evidence_root / 'purchasing/raw/hmda_2025_purchase_oc.csv')
    fields = ['county_code', 'census_tract', 'action_taken', 'loan_purpose', 'occupancy_type',
              'lien_status', 'construction_method', 'total_units', 'reverse_mortgage',
              'open-end_line_of_credit', 'business_or_commercial_purpose',
              'income', 'property_value', 'loan_to_value_ratio']
    db.execute('CREATE TABLE loans (' + ','.join('"'+f+'" TEXT' for f in fields) + ')')
    db.executemany('INSERT INTO loans VALUES ('+','.join('?' for _ in fields)+')',
                   [[r.get(f) for f in fields] for r in raw])
    counties = db.execute('SELECT DISTINCT county_code FROM loans').fetchall()
    assert counties == [('06059',)], counties
    db.create_function('finite_number', 1, num)
    core = [r['tract_geoid'] for r in rows(ROOT/'data/clean/hmda_tract_city_crosswalk.csv')
            if r['city'] == 'Irvine' and r['core_99999'] == 'True']
    db.execute('CREATE TABLE core (tract TEXT PRIMARY KEY)')
    db.executemany('INSERT INTO core VALUES (?)', [(c,) for c in core])
    db.execute('''CREATE VIEW strict AS SELECT * FROM loans WHERE county_code='06059'
        AND action_taken='1' AND loan_purpose='1' AND occupancy_type='1' AND lien_status='1'
        AND construction_method='1' AND total_units='1' AND reverse_mortgage='2'
        AND "open-end_line_of_credit"='2' AND business_or_commercial_purpose='2' ''')
    db.execute('CREATE VIEW irvine AS SELECT * FROM strict JOIN core ON census_tract=tract')
    db.execute('''CREATE VIEW band AS SELECT * FROM irvine
                  WHERE finite_number(property_value)>=2000000 AND finite_number(property_value)<3000000''')
    def median(table, field):
        vals = [x[0] for x in db.execute(f'SELECT finite_number({field}) AS x FROM {table} WHERE x IS NOT NULL ORDER BY x')]
        n = len(vals)
        return {'valid': n, 'median': (vals[(n-1)//2]+vals[n//2])/2,
                'nonpositive': sum(x <= 0 for x in vals)}
    hmda = {'raw_count': len(raw), 'core_tracts': len(core),
            'oc_strict': db.execute('SELECT count(*) FROM strict').fetchone()[0]}
    for table in ['irvine', 'band']:
        hmda[table] = {'loans': db.execute('SELECT count(*) FROM '+table).fetchone()[0],
                       'income_thousands_usd': median(table, 'income'),
                       'cltv_pct': median(table, 'loan_to_value_ratio')}
    assert (hmda['oc_strict'], hmda['irvine']['loans'], hmda['irvine']['income_thousands_usd']['valid'],
            hmda['irvine']['income_thousands_usd']['median']) == (12367, 1002, 915, 326)
    assert (hmda['band']['loans'], hmda['band']['income_thousands_usd']['valid'],
            hmda['band']['income_thousands_usd']['median'], hmda['band']['cltv_pct']['valid'],
            hmda['band']['cltv_pct']['median']) == (159, 139, 463, 156, 69.816)
    out['hmda_sql'] = hmda
    from openpyxl import load_workbook
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        book = load_workbook(args.evidence_root/'supply/raw/irvine_apr2025.xlsm', read_only=True, data_only=True)
        a2 = list(book['Table A2'].iter_rows(min_row=13, values_only=True))
        stages = {}
        for label, ci, cell, expected in [('entitled',19,'C29',4835),('permitted',32,'D29',1411),('completed',45,'E29',1052)]:
            by_type = Counter()
            for r in a2:
                if r[5] and isinstance(r[ci], (int, float)):
                    by_type[str(r[5])] += r[ci]
            assert sum(by_type.values()) == book['Summary'][cell].value == expected
            stages[label] = dict(by_type)
        book.close()
    assert stages['permitted']['SFD'] == 360 and stages['completed']['SFD'] == 600
    out['apr_readonly_workbook'] = stages
    prices = {}
    for r in rows(args.evidence_root/'pricing2/raw/Neighborhood_zhvi_bdrmcnt_4_uc_sfrcondo_tier_0.33_0.67_sm_sa_month.csv'):
        if r['RegionID'] not in {'775702','403127','403197','403196'}:
            continue
        assert r['City'] == 'Irvine' and r['State'] == 'CA'
        last = max(k for k,v in r.items() if k[:4].isdigit() and v)
        assert last == '2026-08-31'
        peak = max(Decimal(v) for k,v in r.items() if '2024-01-01' <= k <= '2025-12-31' and v)
        prices[r['RegionName']] = {'id':r['RegionID'],'yoy_pct':pct(r[last],r['2025-08-31']),
                                  'two_year_pct':pct(r[last],r['2024-08-31']),
                                  'from_2024_2025_peak_pct':pct(r[last],peak)}
    assert len(prices) == 4
    out['zhvi_decimal_ratios'] = prices
    county = {r['month_date_yyyymm']:r for r in rows(args.evidence_root/'pricing3/raw/realtor_county_history.csv')
              if r['county_fips'].zfill(5) == '06059'}
    assert max(county) == '202609'
    result = {}
    for f in ['active_listing_count','pending_listing_count','new_listing_count','median_days_on_market',
              'median_listing_price','median_square_feet','median_listing_price_per_square_foot']:
        result[f] = {'yoy_pct': pct(county['202609'][f],county['202509'][f]),
                     'reported_yoy_pct':float(Decimal(county['202609'][f+'_yy'])*100)}
    result['active_vs_2019_pct'] = pct(county['202609']['active_listing_count'],county['201909']['active_listing_count'])
    out['realtor_decimal_ratios'] = result
    # This additional statement did not belong to the old four-family audit.
    r0, r1 = Decimal('7.375')/1200, Decimal('6.5')/1200
    factor = lambda r:r/(1-(1+r)**(-360))
    out['new_claim_price_plus_10_fixed_cash_PI_change_pct'] = float((70*factor(r1)/(60*factor(r0))-1)*100)
    assert round(out['new_claim_price_plus_10_fixed_cash_PI_change_pct'],2) == 6.77
    params = json.loads((ROOT/'data/step12_a_parameter_interface.json').read_text())
    assert len(params['parameters']) == len({p['id'] for p in params['parameters']}) == 30
    assert params['eligible_for_automatic_A_import'] is False
    counts = Counter()
    for p in params['parameters']:
        counts[p['adoption']] += 1
        assert p['adoption'] in {'reference_only','blocked_pending_evidence'}
        if p['evidence_kind'] == 'missing':
            assert p['base_value'] is None and p['stress_value'] is None
            assert p['adoption'] == 'blocked_pending_evidence'
    out['all_parameter_adoption_counts'] = dict(counts)
    out['source_discrepancies_retained'] = ['Realtor DOM and area reported yoy vs frozen levels; reason unknown']
    out['limitations'] = ['Geography crosswalk reused, no new spatial audit',
                          'No target matched transactions, signed leases, property expenses or personal quotes acquired',
                          'No timing return or household choice inferred from these checks']
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:out[k] for k in ['hmda_sql','apr_readonly_workbook','zhvi_decimal_ratios',
                                      'realtor_decimal_ratios','new_claim_price_plus_10_fixed_cash_PI_change_pct',
                                      'all_parameter_adoption_counts']},ensure_ascii=False))


if __name__ == '__main__':
    main()
