"""Rebuild DOF/Census county components; keep vintages and residuals separate."""
import argparse
import csv
import hashlib
import json
from pathlib import Path
from urllib.request import urlopen
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCES = {
    'dof_e2.xlsx': ('https://dof.ca.gov/media/docs/forecasting/Demographics/estimates/E-2/E-2_Report_July_2020_2025_Feb26_w.xlsx', '967bad41ec264a6e9e597af810ca48226bdf78838485f1f396cebdd5eaf9b61d'),
    'dof_e6.xlsx': ('https://dof.ca.gov/media/docs/forecasting/Demographics/estimates/E-6/E-6_Report_July_2020-2025_Feb26_w.xlsx', 'a6797dfcae13cc89bb078c289148b83fc68c37788c3b0b34ea50144c36180cff'),
    'co-est2025-alldata.csv': ('https://www2.census.gov/programs-surveys/popest/datasets/2020-2025/counties/totals/co-est2025-alldata.csv', '4f5a499d851e2cb48fd7a5405e5a9235453a8a66933657aacd10df0e264f35d5')
}
GEOS = {'California': 0, 'Los Angeles': 37, 'Orange': 59, 'Riverside': 65, 'San Bernardino': 71, 'San Diego': 73}
MEASURES = ['population_start', 'population_end', 'population_change', 'births', 'deaths', 'natural_change', 'net_migration', 'net_international_migration', 'net_domestic_migration', 'residual']


def name(g):
    return g if g == 'California' else g + ' County'


def record(g, year, source, vintage, locator, values):
    return dict(geography=name(g), state_fips='06', county_fips=f'{GEOS[g]:03d}',
                geography_type='state' if g == 'California' else 'county',
                period_start=f'{year-1}-07-01', period_end_exclusive=f'{year}-07-01',
                estimate_year=year, source_id=source, publication_vintage=vintage,
                source_locator=locator, **dict(zip(MEASURES, values)))


def build(raw_dir):
    fingerprints = {}
    for file, (_, expected) in SOURCES.items():
        digest = hashlib.sha256((raw_dir/file).read_bytes()).hexdigest()
        if digest != expected:
            raise ValueError(f'Changed fingerprint: {file}; inspect revised vintage first.')
        fingerprints[file] = digest
    e2 = load_workbook(raw_dir/'dof_e2.xlsx', data_only=True, read_only=True)
    e6 = load_workbook(raw_dir/'dof_e6.xlsx', data_only=True, read_only=True)
    dof = []
    for year in range(2021, 2026):
        sheet = e2[f'July {year-1}-{year} components']
        for row_num, r in enumerate(sheet.iter_rows(values_only=True), 1):
            g = str(r[0]).strip()
            if g not in GEOS:
                continue
            values = [r[1], r[2], r[3], r[5], r[6], r[7], r[8], r[9], r[10], 0]
            dof.append(record(g, year, 'SRC-012', '2026-02-04', f'{sheet.title}!row{row_num}', values))
    e6_index = {}
    current = None
    for r in e6['E-6 2020-2025 Report'].iter_rows(values_only=True):
        if r[0] is not None:
            label = str(r[0]).strip()
            # This pinned workbook splits these county names across adjacent rows.
            current = {'Los': 'Los Angeles', 'Angeles': 'Los Angeles',
                       'Bernardino': 'San Bernardino'}.get(label, label)
        if current in GEOS and str(r[1]) in {str(y) for y in range(2021, 2026)}:
            e6_index[(name(current), int(r[1]))] = [r[2], r[4], *r[5:11]]
    failures = []
    for r in dof:
        values = [r[k] for k in ['population_end','population_change','births','deaths','natural_change','net_migration','net_international_migration','net_domestic_migration']]
        if e6_index.get((r['geography'],r['estimate_year'])) != values:
            failures.append(f'E2/E6 mismatch {r["geography"]}/{r["estimate_year"]}')
    with (raw_dir/'co-est2025-alldata.csv').open(encoding='cp1252', newline='') as f:
        raw = list(csv.DictReader(f))
    census = []
    for g, county in GEOS.items():
        found = [r for r in raw if int(r['STATE']) == 6 and int(r['COUNTY']) == county and int(r['SUMLEV']) == (40 if county == 0 else 50)]
        if len(found) != 1:
            raise ValueError(f'Ambiguous Census geography {g}')
        r = found[0]
        if r['STNAME'] != 'California' or r['CTYNAME'] != name(g):
            raise ValueError(f'Census name/FIPS mismatch {g}: {r["CTYNAME"]}')
        for year in range(2021,2026):
            y=str(year)
            values = [int(r['POPESTIMATE'+str(year-1)]), int(r['POPESTIMATE'+y]),
                      *[int(r[k+y]) for k in ['NPOPCHG','BIRTHS','DEATHS','NATURALCHG','NETMIG','INTERNATIONALMIG','DOMESTICMIG','RESIDUAL']]]
            census.append(record(g, year, 'SRC-015', 'Vintage 2025; released 2026-03', f'STATE06/COUNTY{county:03d}/SUMLEV{r["SUMLEV"]}', values))
    records = dof+census
    for r in records:
        prefix=f'{r["source_id"]}/{r["geography"]}/{r["estimate_year"]}'
        if any(r[m] is None for m in MEASURES):
            failures.append(prefix+': missing component')
        if r['population_end']-r['population_start'] != r['population_change']:
            failures.append(prefix+': population delta')
        if r['births']-r['deaths'] != r['natural_change']:
            failures.append(prefix+': natural change')
        if r['net_domestic_migration']+r['net_international_migration'] != r['net_migration']:
            failures.append(prefix+': migration sum')
        if r['natural_change']+r['net_migration']+r['residual'] != r['population_change']:
            failures.append(prefix+': residual accounting')
    keys={(r['source_id'],r['geography'],r['estimate_year']) for r in records}
    if len(records)!=60 or len(keys)!=60:
        failures.append('Expected 60 unique source/geography/year records')
    records.sort(key=lambda r:(r['source_id'],r['geography'],r['estimate_year']))
    qa={'status':'pass' if not failures else 'fail','rows':len(records),'unique_keys':len(keys),
        'dof_e2_e6_matched_rows':len(dof),'fingerprints':fingerprints,'failures':failures,
        'limits':['E2 and E6 are two presentations of one DOF series, not independent evidence.',
                  'DOF and Census share some administrative inputs; agreement is not independent replication.',
                  'Each 2021–2025 record covers July 1 of prior year through June 30 of named year.',
                  '2020 Apr–Jun components excluded from full-year/cumulative comparisons.',
                  'County migration does not identify Irvine community migration or target home buyers.',
                  'Census residual is retained; DOF components sum exactly in this extraction.']}
    (ROOT/'qa').mkdir(exist_ok=True)
    (ROOT/'qa/demographic_components_qa.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2)+'\n')
    if failures:
        raise ValueError(failures)
    with (ROOT/'data/clean/county_components_2020_2025.csv').open('w',encoding='utf-8',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(records[0]),lineterminator='\n');w.writeheader();w.writerows(records)
    comparisons=[]
    for year in range(2021,2026):
        a=next(r for r in dof if r['geography']=='Orange County' and r['estimate_year']==year)
        b=next(r for r in census if r['geography']=='Orange County' and r['estimate_year']==year)
        comparisons.append({'estimate_year':year,**{f'dof_{m}':a[m] for m in MEASURES},
                            **{f'census_{m}':b[m] for m in MEASURES},
                            'same_net_migration_sign':(a['net_migration']>0)==(b['net_migration']>0)})
    with (ROOT/'data/clean/orange_county_vintage_comparison.csv').open('w',encoding='utf-8',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(comparisons[0]),lineterminator='\n');w.writeheader();w.writerows(comparisons)
    print(json.dumps({'status':qa['status'],'rows':len(records),'dof_e2_e6_matched_rows':len(dof),'orange_county_same_sign_years':sum(r['same_net_migration_sign'] for r in comparisons)}))


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--raw-dir',type=Path,default=ROOT/'data/raw');p.add_argument('--download',action='store_true');a=p.parse_args()
    if a.download:
        a.raw_dir.mkdir(parents=True,exist_ok=True)
        for file,(url,_) in SOURCES.items():
            with urlopen(url,timeout=40) as r:b=r.read()
            if (file.endswith('.xlsx') and not b.startswith(b'PK')) or b.lstrip().startswith(b'<'):
                raise ValueError(f'Unexpected error/login response for {file}')
            (a.raw_dir/file).write_bytes(b)
    build(a.raw_dir)
